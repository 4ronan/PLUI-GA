import json, subprocess, zipfile
from pathlib import Path
from openpyxl import load_workbook

URL='https://www.insee.fr/fr/statistiques/fichier/8998387/TD_LOGEMENT_CARACT_PRINC_2023_xlsx.zip'
TARGET='16015'
OUT=Path('output/insee-log1-3kc-audit.json')
TMP=Path('tmp_3kc')
TMP.mkdir(exist_ok=True)
zip_path=TMP/'TD_LOGEMENT_CARACT_PRINC_2023_xlsx.zip'

subprocess.run([
    'curl','--http1.1','--fail','--location','--show-error','--silent',
    '--retry','5','--retry-all-errors','--retry-delay','2',
    '--connect-timeout','30','--max-time','900',
    '--output',str(zip_path),URL
], check=True)

if not zipfile.is_zipfile(zip_path):
    raise RuntimeError('La ressource INSEE XLSX n’est pas une archive ZIP valide')

with zipfile.ZipFile(zip_path) as z:
    xlsx_names=[n for n in z.namelist() if n.lower().endswith('.xlsx')]
    if not xlsx_names:
        raise RuntimeError('Aucun fichier XLSX dans l’archive INSEE')
    xlsx_name=max(xlsx_names,key=lambda n:z.getinfo(n).file_size)
    xlsx_path=TMP/Path(xlsx_name).name
    xlsx_path.write_bytes(z.read(xlsx_name))

wb=load_workbook(xlsx_path,read_only=True,data_only=True)

sheet_audit=[]
target_sheet=None
headers=None
target_values=None
header_row_num=None

# Le classeur XLSX diffusé par l'Insee n'utilise pas nécessairement les noms
# techniques Melodi (CODGEO/GEO). On repère d'abord la ligne de la commune,
# puis on remonte vers la ligne d'en-tête la plus plausible.
for ws in wb.worksheets:
    preview=[]
    for row_idx,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,12),values_only=True),start=1):
        preview.append({
            'row':row_idx,
            'values':[None if v is None else str(v) for v in row[:12]]
        })

    explicit_header=None
    explicit_header_row=None
    for row_idx,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,30),values_only=True),start=1):
        vals=[str(v).strip() if v is not None else '' for v in row]
        normalized=[v.upper().replace(' ','').replace('_','') for v in vals]
        if any(v in {'CODGEO','GEO','CODGÉO','CODEGÉOGRAPHIQUE','CODEGEOGRAPHIQUE'} for v in normalized):
            explicit_header=vals
            explicit_header_row=row_idx
            break

    target_row_num=None
    target_row=None
    # La commune doit se trouver dans les premières colonnes ; on limite à 8
    # pour éviter qu'une valeur indicatrice égale à 16015 soit prise pour un code.
    for row_idx,row in enumerate(ws.iter_rows(values_only=True),start=1):
        first=[str(v).strip() if v is not None else '' for v in row[:8]]
        if TARGET in first:
            target_row_num=row_idx
            target_row=list(row)
            break

    chosen_header=None
    chosen_header_row=None
    if target_row_num is not None:
        if explicit_header is not None and explicit_header_row < target_row_num:
            chosen_header=explicit_header
            chosen_header_row=explicit_header_row
        else:
            # Cherche, dans les 12 lignes précédentes, la ligne contenant le plus
            # de libellés textuels/non vides. Les tableaux Insee placent souvent
            # l'en-tête juste au-dessus des données avec 1 à 3 lignes de titre.
            start=max(1,target_row_num-12)
            candidates=[]
            for rnum,row in enumerate(ws.iter_rows(min_row=start,max_row=target_row_num-1,values_only=True),start=start):
                vals=[str(v).strip() if v is not None else '' for v in row]
                nonempty=sum(bool(v) for v in vals)
                textual=sum(bool(v) and not v.replace('.','',1).replace(',','',1).isdigit() for v in vals)
                candidates.append((textual,nonempty,rnum,vals))
            if candidates:
                textual,nonempty,chosen_header_row,chosen_header=max(candidates,key=lambda x:(x[0],x[1],x[2]))

    sheet_audit.append({
        'sheet':ws.title,
        'max_row':ws.max_row,
        'max_column':ws.max_column,
        'explicit_header_row':explicit_header_row,
        'target_row':target_row_num,
        'chosen_header_row':chosen_header_row,
        'preview':preview
    })

    if target_row is not None and chosen_header is not None:
        target_sheet=ws.title
        headers=chosen_header
        target_values=target_row
        header_row_num=chosen_header_row
        break

if target_values is None:
    raise RuntimeError(f'Aucune ligne trouvée pour {TARGET}; feuilles={sheet_audit}')

# Garantit des noms de colonnes uniques et exploitables, même lorsque l'en-tête
# contient des cellules vides ou fusionnées.
clean_headers=[]
seen={}
for i,h in enumerate(headers):
    base=(str(h).strip() if h is not None else '') or f'COL_{i+1}'
    n=seen.get(base,0)+1
    seen[base]=n
    clean_headers.append(base if n==1 else f'{base}__{n}')
headers=clean_headers

record={}
for i,h in enumerate(headers):
    record[h]=target_values[i] if i < len(target_values) else None

# Audit sémantique par noms de colonnes : aucune interprétation n'est encore figée ici.
def cols_matching(tokens):
    return [h for h in headers if h and any(t.lower() in h.lower() for t in tokens)]

vacancy_cols=cols_matching(['VAC','DW_VAC'])
construction_cols=cols_matching(['LT1919','1919','1946','1971','1991','2006','ACH','PERACH','CONSTR'])
dwelling_type_cols=cols_matching(['MAISON','APPART','HOUSE','APART','TYPL','TYPE'])
nor_cols=cols_matching(['NOR'])

result={
    'source':'Insee RP2023 - LOG1 / DS_RP_TD_LOGEMENT_CARACT_PRINC',
    'url':URL,
    'territory':TARGET,
    'format':'xlsx_zip',
    'xlsx_file':xlsx_name,
    'sheet':target_sheet,
    'header_row':header_row_num,
    'fields_n':len(headers),
    'fields':headers,
    'target_rows':1,
    'target_rows_2023':1,
    'target_record':record,
    'sheet_audit':sheet_audit,
    'detected':{
        'vacancy_columns':vacancy_cols,
        'construction_period_columns':construction_cols,
        'dwelling_type_columns':dwelling_type_cols,
        'nor_columns_excluded_from_vacant_profile':nor_cols
    },
    'quality':{
        'year':2023,
        'geography':'commune',
        'note':'Audit du tableau détaillé LOG1 officiel; aucune fusion avec LOVAC et aucun calcul causal.',
        'nor_rule':'NOR appartient aux résidences principales et ne doit pas être utilisé pour le profil des logements vacants.'
    }
}

OUT.parent.mkdir(exist_ok=True)
OUT.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2,default=str))
